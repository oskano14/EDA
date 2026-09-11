from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Lumina_Co_Reactivation_Clients_a_risque_1.pptx"
OUTPUT = ROOT / "Lumina_Co_Restitution_CMO_complete.pptx"
CHART = ROOT / "attribution_multicanal.png"

INK = RGBColor(61, 27, 61)
PLUM = RGBColor(92, 58, 92)
RED = RGBColor(201, 24, 74)
PINK = RGBColor(232, 180, 188)
PALE = RGBColor(247, 232, 234)
MUTED = RGBColor(138, 110, 130)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(23, 126, 137)


def set_background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(
    slide,
    text,
    x,
    y,
    width,
    height,
    size=20,
    color=INK,
    font="Calibri",
    bold=False,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_header(slide, eyebrow, title, subtitle=None, dark=False):
    primary = WHITE if dark else INK
    secondary = PINK if dark else RED
    muted = PINK if dark else MUTED
    add_text(slide, eyebrow.upper(), 0.7, 0.35, 11.9, 0.3, 9, secondary, bold=True)
    add_text(slide, title, 0.7, 0.78, 11.9, 0.65, 25, primary, "Cambria", True)
    if subtitle:
        add_text(slide, subtitle, 0.7, 1.38, 11.9, 0.45, 12, muted)


def add_footer(slide, number, dark=False):
    color = PINK if dark else MUTED
    add_text(slide, "Lumina & Co  ·  Restitution au CMO", 0.7, 7.15, 5.5, 0.2, 8, color)
    add_text(slide, str(number), 12.15, 7.15, 0.45, 0.2, 8, color, align=PP_ALIGN.RIGHT)


def add_metric(slide, x, y, width, value, label, accent=RED):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(1.28),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = PALE
    add_text(slide, value, x + 0.22, y + 0.16, width - 0.44, 0.45, 24, accent, "Cambria", True)
    add_text(slide, label, x + 0.22, y + 0.70, width - 0.44, 0.42, 10, INK)


def add_action(slide, number, title, detail, y):
    marker = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(0.8), Inches(y), Inches(0.48), Inches(0.48)
    )
    marker.fill.solid()
    marker.fill.fore_color.rgb = RED
    marker.line.color.rgb = RED
    frame = marker.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = str(number)
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = WHITE
    add_text(slide, title, 1.5, y - 0.02, 4.0, 0.35, 15, INK, "Cambria", True)
    add_text(slide, detail, 5.0, y - 0.02, 7.4, 0.62, 11, MUTED)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_kpi_slide(prs, layout):
    slide = prs.slides.add_slide(layout)
    set_background(slide, WHITE)
    add_header(
        slide,
        "KPIs de campagne",
        "La campagne gagnante dépend de l'objectif",
        "Winter Promo maximise le rendement ; Spring Launch recrute au meilleur coût.",
    )
    add_metric(slide, 0.8, 2.05, 2.7, "5,35", "ROAS · Winter Promo")
    add_metric(slide, 3.72, 2.05, 2.7, "9,45 €", "CPA · Winter Promo")
    add_metric(slide, 6.64, 2.05, 2.7, "573 €", "CAC · Spring Launch", TEAL)
    add_metric(slide, 9.56, 2.05, 2.7, "76,5 %", "Conversion · Spring Launch", TEAL)
    add_text(slide, "Le classement change avec le KPI", 0.8, 3.75, 5.2, 0.4, 18, INK, "Cambria", True)
    add_text(
        slide,
        "Un CPA bas récompense les réachats. Le CAC isole les primo-achats. Le ROAS rapporte le revenu au coût. Aucun indicateur ne suffit seul.",
        0.8,
        4.25,
        5.35,
        1.2,
        13,
        MUTED,
    )
    warning = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.55), Inches(3.72), Inches(5.7), Inches(1.75)
    )
    warning.fill.solid()
    warning.fill.fore_color.rgb = INK
    warning.line.color.rgb = INK
    add_text(slide, "POINT DE VIGILANCE", 6.85, 4.0, 4.9, 0.25, 9, PINK, bold=True)
    add_text(
        slide,
        "Summer Sale ferme le classement : ROAS 2,74 et CAC 2 025 €. À investiguer avant reconduction.",
        6.85,
        4.36,
        4.95,
        0.75,
        14,
        WHITE,
        "Cambria",
        True,
    )
    add_footer(slide, 4)
    add_notes(
        slide,
        "[~1 min] Les KPI ne désignent pas tous le même gagnant. Winter Promo a le meilleur ROAS, 5,35, et le CPA le plus bas, 9,45 euros. Spring Launch a le meilleur taux de conversion, 76,5 %, et le CAC le plus faible, 573 euros. Le choix dépend donc de l'objectif : réachat, acquisition ou revenu. Summer Sale est le signal faible à investiguer avant reconduction.",
    )
    return slide


def add_attribution_slide(prs, layout):
    slide = prs.slides.add_slide(layout)
    set_background(slide, WHITE)
    add_header(
        slide,
        "Attribution multicanal",
        "Le last touch efface 82 % du rôle de découverte",
        "94 366 conversions · 5,03 contacts par parcours · 99 % de parcours multicanaux",
    )
    slide.shapes.add_picture(str(CHART), Inches(0.62), Inches(1.85), width=Inches(8.35))
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.15), Inches(1.9), Inches(3.45), Inches(4.45)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = INK
    panel.line.color.rgb = INK
    add_text(slide, "FIRST TOUCH", 9.48, 2.25, 2.8, 0.25, 9, PINK, bold=True)
    add_text(slide, "82,0 %", 9.48, 2.58, 2.8, 0.5, 27, WHITE, "Cambria", True)
    add_text(slide, "du CA commence par display ou social", 9.48, 3.12, 2.7, 0.65, 11, PINK)
    add_text(slide, "LAST TOUCH", 9.48, 4.05, 2.8, 0.25, 9, PINK, bold=True)
    add_text(slide, "53,1 %", 9.48, 4.38, 2.8, 0.5, 27, WHITE, "Cambria", True)
    add_text(slide, "du CA finit par affiliation ou retargeting", 9.48, 4.92, 2.7, 0.65, 11, PINK)
    add_text(slide, "Conclusion : aucun modèle mono-touch ne suffit.", 9.48, 5.72, 2.7, 0.4, 11, WHITE, bold=True)
    add_footer(slide, 5)
    add_notes(
        slide,
        "[~1 min 15] Le parcours moyen comporte cinq contacts et 99 % des parcours convertis sont multicanaux. En first touch, display et social ouvrent 82 % du chiffre d'affaires. En last touch, ils disparaissent entièrement, tandis que l'affiliation et le retargeting captent 53,1 %. Une lecture uniquement au dernier clic ferait couper les canaux qui créent la découverte.",
    )
    return slide


def add_budget_slide(prs, layout):
    slide = prs.slides.add_slide(layout)
    set_background(slide, WHITE)
    add_header(
        slide,
        "Arbitrage budgétaire",
        "Rééquilibrer par le test, pas par une règle d'attribution",
        "L'attribution décrit le parcours ; seule l'incrémentalité mesure l'effet réel.",
    )
    add_action(slide, 1, "Auditer l'affiliation", "83,7 % du coût observé, contre 15,7 % du CA en attribution linéaire.", 2.05)
    add_action(slide, 2, "Protéger la découverte", "Conserver un budget test display et social : ils ouvrent 82,0 % du CA attribué.", 3.15)
    add_action(slide, 3, "Tester les canaux de fermeture", "Holdouts pour email et retargeting ; suivi CPA, CAC, marge et revenu incrémental.", 4.25)
    warning = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.55), Inches(11.45), Inches(0.92)
    )
    warning.fill.solid()
    warning.fill.fore_color.rgb = PALE
    warning.line.color.rgb = PALE
    add_text(slide, "CONFIANCE DATA", 1.05, 5.82, 1.45, 0.22, 9, RED, bold=True)
    add_text(
        slide,
        "Réconcilier le CA campagnes (5,56 M€) et le CA des factures converties (18,40 M€) avant une réallocation majeure.",
        2.55,
        5.72,
        9.25,
        0.45,
        12,
        INK,
        bold=True,
    )
    add_footer(slide, 6)
    add_notes(
        slide,
        "[~1 min] Notre recommandation n'est pas de déplacer mécaniquement le budget selon une colonne. D'abord, auditer l'affiliation, qui concentre 83,7 % des coûts. Ensuite, protéger un budget test pour display et social. Enfin, mesurer email et retargeting avec des groupes témoins. Avant tout arbitrage majeur, il faut réconcilier la définition du revenu entre campaigns.csv et les factures.",
    )
    return slide


def replace_text(slide, replacements):
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                for old, new in replacements.items():
                    if old in run.text:
                        run.text = run.text.replace(old, new)


def move_slide(prs, slide, index):
    slide_ids = prs.slides._sldIdLst
    slide_id_element = next(
        element for element in slide_ids if int(element.id) == slide.slide_id
    )
    slide_ids.remove(slide_id_element)
    slide_ids.insert(index, slide_id_element)


def main():
    if not SOURCE.exists() or not CHART.exists():
        raise FileNotFoundError("La présentation source ou le graphique d'attribution est absent.")

    prs = Presentation(SOURCE)
    layout = min(prs.slide_layouts, key=lambda candidate: len(candidate.placeholders))
    kpi_slide = add_kpi_slide(prs, layout)
    attribution_slide = add_attribution_slide(prs, layout)
    budget_slide = add_budget_slide(prs, layout)

    move_slide(prs, kpi_slide, 3)
    move_slide(prs, attribution_slide, 4)
    move_slide(prs, budget_slide, 5)

    replace_text(
        prs.slides[0],
        {"SEGMENTATION RFM": "RESTITUTION CMO"},
    )
    replace_text(
        prs.slides[7],
        {
            "≈ 167 000 €": "≈ 32 000 €",
            "si 10 % des 4 704 clients à risque se réactivent, à leur valeur médiane historique (355 €).": "pour un gain incrémental de 5 points, soit environ 235 clients à 137 € par commande.",
        },
    )
    add_notes(
        prs.slides[7],
        "[~45 sec] Un gain incrémental de cinq points représenterait environ 235 clients réactivés et 32 000 euros de chiffre d'affaires brut sur la base d'une commande moyenne historique de 137 euros. Ce n'est pas une prévision : le groupe témoin doit mesurer le gain réel. La décision demandée est de valider le pilote, son protocole et son seuil de succès.",
    )

    prs.save(OUTPUT)
    print(f"Présentation créée : {OUTPUT.name} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()